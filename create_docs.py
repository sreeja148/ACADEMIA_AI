import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

doc = docx.Document()

# Page Setup
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

# Color Palette
PRIMARY_COLOR = RGBColor(79, 70, 229)    # #4f46e5 (Indigo)
SECONDARY_COLOR = RGBColor(6, 182, 212)  # #06b6d4 (Cyan)
DARK_TEXT = RGBColor(15, 23, 42)        # #0f172a
MUTED_TEXT = RGBColor(100, 116, 139)    # #64748b

# Base Style
normal_style = doc.styles['Normal']
normal_style.font.name = 'Calibri'
normal_style.font.size = Pt(11)
normal_style.font.color.rgb = DARK_TEXT

# Title
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_p.paragraph_format.space_after = Pt(4)
title_run = title_p.add_run("ACADEMIA AI")
title_run.font.name = 'Calibri'
title_run.font.size = Pt(28)
title_run.font.bold = True
title_run.font.color.rgb = PRIMARY_COLOR

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_p.paragraph_format.space_after = Pt(24)
sub_run = sub_p.add_run("AI-Powered Professor Recommendation & Smart Course Planning System")
sub_run.font.name = 'Calibri'
sub_run.font.size = Pt(14)
sub_run.font.italic = True
sub_run.font.color.rgb = SECONDARY_COLOR

# Divider
p_div = doc.add_paragraph()
p_div.paragraph_format.space_after = Pt(18)
r_div = p_div.add_run("_________________________________________________________________________________")
r_div.font.color.rgb = RGBColor(226, 232, 240)

def add_heading_1(text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(18)
    h.paragraph_format.space_after = Pt(8)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = PRIMARY_COLOR

def add_heading_2(text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = SECONDARY_COLOR

def add_callout(text, bold_prefix=""):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "EEF2FF")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if bold_prefix:
        r_bold = p.add_run(bold_prefix + " ")
        r_bold.bold = True
        r_bold.font.color.rgb = PRIMARY_COLOR
    r_text = p.add_run(text)
    r_text.font.size = Pt(10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

# Section 1: Project Overview
add_heading_1("1. Executive Summary & Project Overview")
p1 = doc.add_paragraph(
    "Academia AI is a state-of-the-art web application designed to solve a fundamental problem faced by higher education students: "
    "selecting the right professors and designing balanced semester schedules. Rather than relying on simple 5-star rating averages, "
    "Academia AI leverages client-side Machine Learning (TensorFlow.js), Natural Language Processing (NLP), and multi-variable suitability models "
    "to match students with faculty members based on teaching quality, grading leniency, difficulty alignment, and individual learning preferences."
)
p1.paragraph_format.space_after = Pt(8)

p2 = doc.add_paragraph(
    "The platform features an interactive glassmorphic web interface supporting three specialized user roles—Students, Professors, and Administrators. "
    "All machine learning models execute directly inside the user's web browser, delivering real-time inferences with zero backend server dependencies."
)
p2.paragraph_format.space_after = Pt(12)

add_callout(
    "Academia AI eliminates guesswork from course registration by transforming qualitative student evaluations into actionable, neural-matched compatibility scores.",
    "Core Value Proposition:"
)

# Section 2: Core Technologies
add_heading_1("2. Technical Architecture & Stack")
doc.add_paragraph("The application is constructed with modern web standards and client-side machine learning technologies:")

tech_table = doc.add_table(rows=6, cols=2)
tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["Component Layer", "Technology Stack & Library"]
for col_idx, text in enumerate(headers):
    cell = tech_table.cell(0, col_idx)
    set_cell_background(cell, "4F46E5")
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)

data = [
    ("Core Frontend Structure", "HTML5 (Semantic Layout), Vanilla JavaScript (ES6+)"),
    ("Machine Learning Engine", "TensorFlow.js (Browser Deep Learning Framework)"),
    ("Natural Language Processing", "Custom NLP Sentiment & Aspect Extraction Pipeline"),
    ("Design & Visuals", "Vanilla CSS3 (Glassmorphism, Dark/Light Themes, Micro-interactions)"),
    ("Icons & Analytics", "Lucide SVG Icon Engine & Chart.js Visualizations")
]

for row_idx, (c1, c2) in enumerate(data, start=1):
    row_cells = tech_table.rows[row_idx].cells
    bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
    set_cell_background(row_cells[0], bg)
    set_cell_background(row_cells[1], bg)
    row_cells[0].paragraphs[0].add_run(c1).bold = True
    row_cells[1].paragraphs[0].add_run(c2)

doc.add_paragraph().paragraph_format.space_after = Pt(12)

# Section 3: Machine Learning Capabilities
add_heading_1("3. Machine Learning & AI Capabilities in Detail")

add_heading_2("A. TensorFlow.js Neural Network Match Engine")
doc.add_paragraph(
    "A Multi-Layer Perceptron (MLP) neural network constructed using TensorFlow.js. It extracts feature vectors representing both "
    "student preferences (learning style, target difficulty, past GPA, grading preference) and professor attributes (teaching quality, "
    "grading fairness, responsiveness, lab quality) to produce a 0-100% Neural Match Score along with confidence bounds and decision insights."
)

add_heading_2("B. NLP Review Sentiment & Aspect Mining")
doc.add_paragraph(
    "Evaluates written course reviews ('pros' and 'cons') for emotional polarity and key topics. The model assigns sentiment classification "
    "badges (Positive, Constructive, Critical) and extracts key aspect tags such as 'Exam Focus', 'Responsive', 'Lab Heavy', and 'Heavy Workload'."
)

add_heading_2("C. Student Workload & Burnout Risk Predictor")
doc.add_paragraph(
    "Calculates regression metrics when courses are staged in the Semester Planner. It forecasts total weekly study hours, predicts term cumulative GPA, "
    "and categorizes burnout risk (Low Risk, Moderate Risk, High Risk) with actionable academic advice."
)

add_heading_2("D. Semantic Intent Search Engine")
doc.add_paragraph(
    "Powers the global search bar using vector similarity mapping. Students can type natural language queries like 'easy electives', "
    "'hands-on lab biology', or 'lenient grading math' to instantly find relevant faculty and courses."
)

# Section 4: Feature Breakdown by User Role
add_heading_1("4. Feature Breakdown by User Role")

add_heading_2("A. Student Portal Features")
bullets_student = [
    ("Dashboard Overview: ", "Displays current GPA, ML projected GPA, term burnout risk index, degree completion ring, and top AI-matched faculty."),
    ("Course Catalog & Discovery: ", "Browse all offered courses with department filters and real-time semantic search."),
    ("AI Recommendations: ", "Adjust learning style (Visual, Hands-on, Theoretical, Discussion) and grading preferences (Lenient, Balanced, Rigorous) to dynamically trigger real-time neural score recalculation."),
    ("Semester Planner: ", "Interactive weekly schedule grid with timetable conflict detection, credit-limit validation, and ML workload forecasting."),
    ("GPA Calculator: ", "Interactive target grade simulator predicting cumulative GPA based on completed course credits."),
    ("Course Evaluations & Reviews: ", "Read student feedback with NLP sentiment badges, pros/cons breakdown, and upvote helpful reviews.")
]
for title, desc in bullets_student:
    p = doc.add_paragraph(style='List Bullet')
    r1 = p.add_run(title)
    r1.bold = True
    r1.font.color.rgb = PRIMARY_COLOR
    p.add_run(desc)

add_heading_2("B. Professor Portal Features")
bullets_prof = [
    ("Faculty Dashboard: ", "View student enrollment numbers, written review count, overall evaluation rating, and radar chart of instructional traits."),
    ("Profile Management: ", "Update office hours schedule, bio, and upload syllabus documents for student download."),
    ("Feedback Analytics: ", "Inspect student reviews and sentiment trends across taught courses.")
]
for title, desc in bullets_prof:
    p = doc.add_paragraph(style='List Bullet')
    r1 = p.add_run(title)
    r1.bold = True
    r1.font.color.rgb = PRIMARY_COLOR
    p.add_run(desc)

add_heading_2("C. Admin Portal Features")
bullets_admin = [
    ("Institutional Command Center: ", "Snapshots of total student population, faculty count, active courses, and average institution rating."),
    ("Roster & Catalog CRUD: ", "Manage student profiles, faculty directories, and course offerings."),
    ("AI Analytics Dashboard: ", "Charts tracking popular courses, predicted completion rates, department performance benchmarks, and registration trends."),
    ("Review Moderation: ", "Flag or remove non-compliant course reviews.")
]
for title, desc in bullets_admin:
    p = doc.add_paragraph(style='List Bullet')
    r1 = p.add_run(title)
    r1.bold = True
    r1.font.color.rgb = PRIMARY_COLOR
    p.add_run(desc)

# Section 5: How to Use
add_heading_1("5. How to Use the Application (Step-by-Step Guide)")

steps = [
    ("Step 1: Sign In & Role Selection", "Open the application. On the Auth Screen, choose a role tab (Student, Professor, or Admin). Click 'Sign In' to enter the live prototype."),
    ("Step 2: Explore AI Recommendations (Student)", "Navigate to 'Recommendations'. Select a course from the dropdown, adjust your learning style and grading preference. Observe how the Neural Recommender instantly recalculates compatibility percentages and explanations."),
    ("Step 3: Plan Your Semester", "Navigate to 'Semester Planner'. Add courses from the catalog. Watch the ML Workload Forecaster update your projected GPA, estimated weekly study hours, and burnout risk level. The interactive timetable will highlight any schedule conflicts in red."),
    ("Step 4: Use Semantic Search", "In the top search bar, type queries like 'easy elective' or 'hands-on lab'. The system automatically applies vector semantic filtering to display matching results."),
    ("Step 5: Toggle Dark / Light Theme", "Click the moon/sun icon in the top header bar to switch between Obsidian Dark mode and Crisp Light mode.")
]

for step_title, step_desc in steps:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(step_title)
    r.bold = True
    r.font.color.rgb = PRIMARY_COLOR
    p_desc = doc.add_paragraph(step_desc)
    p_desc.paragraph_format.space_after = Pt(8)

# Section 6: Local Execution & Deployment
add_heading_1("6. Local Execution & GitHub Vercel Deployment")
doc.add_paragraph("To run the project locally:")
doc.add_paragraph("1. Open PowerShell or Command Prompt in the project folder.\n2. Run: python -m http.server 8089\n3. Open http://localhost:8089 in any web browser.")

doc.add_paragraph("To deploy on Vercel:\n1. Push the code to the GitHub repository: https://github.com/abubakarsiddiq3215a/LSE_Prototype\n2. Import the repository into Vercel as a Static Site.\n3. Vercel will instantly host index.html with full client-side TensorFlow.js support.")

output_path = r"c:\ABUBAKAR\B-TECH\ABUBAKAR\SREEJA\LSE - PROTOTYPE\academia\Academia_AI_Project_Documentation.docx"
doc.save(output_path)
print("Successfully generated Word document at:", output_path)
